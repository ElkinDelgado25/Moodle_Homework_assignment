"""Reconstruye la portada editable; requiere python-docx solo para desarrollo."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ASSETS = Path(__file__).resolve().parents[1] / "src/moodle_tasks/assets"


def build_cover(output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin, section.bottom_margin = Pt(82), Cm(2.5)
    section.header_distance = Cm(0.8)
    section.different_first_page_header_footer = True
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    for name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in style._element.xpath(".//w:pBdr"):
            border.getparent().remove(border)

    def page_number(paragraph):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        run = OxmlElement("w:r")
        properties = OxmlElement("w:rPr")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "18")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "808080")
        properties.extend([size, color])
        run.append(properties)
        text = OxmlElement("w:t")
        text.text = "1"
        run.append(text)
        field.append(run)
        paragraph._p.append(field)

    regular_header = section.header.paragraphs[0]
    regular_header.paragraph_format.right_indent = Cm(-1.8)
    page_number(regular_header)
    first_header = section.first_page_header
    first_header.paragraphs[0].paragraph_format.line_spacing = Pt(1)
    table = first_header.add_table(rows=1, cols=2, width=Cm(19.6))
    table.autofit = False
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), str(int(Cm(-1.8).twips)))
    indent.set(qn("w:type"), "dxa")
    table._tbl.tblPr.append(indent)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    table.cell(0, 0).paragraphs[0].add_run().add_picture(str(ASSETS / "uleam-logo.png"), width=Cm(3.75))
    page_number(table.cell(0, 1).paragraphs[0])

    lines = [
        ("Universidad Laica “Eloy Alfaro de Manabí”", True, 82),
        ("Materia:", True, 128),
        ("{{subject}}", False, 151),
        ("Docente:", True, 175),
        ("{{teacher}}", True, 199),
        ("Estudiantes:", True, 246),
        ("{{student_1}}", False, 270),
        ("{{student_2}}", False, 294),
        ("Carrera:", True, 342),
        ("{{degree}}", False, 366),
        ("Curso:", True, 414),
        ("{{class_group}}", False, 438),
        ("Año:", True, 461),
        ("{{year}}", False, 486),
    ]
    for index, (text, bold, position) in enumerate(lines):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.line_spacing = Pt(13)
        paragraph.paragraph_format.keep_with_next = False
        if index + 1 < len(lines):
            paragraph.paragraph_format.space_after = Pt(lines[index + 1][2] - position - 13)
        paragraph.add_run(text).bold = bold
    document.core_properties.title = "Plantilla de portada académica ULEAM"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


if __name__ == "__main__":
    build_cover(ASSETS / "academic-cover.docx")
